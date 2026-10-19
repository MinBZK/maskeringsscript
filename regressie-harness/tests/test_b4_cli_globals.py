"""Regression checks for the resolved CLI path and output-column bugs.

A fake CombinedAnonymizer avoids loading a model. These are passing regression
checks; the outstanding strict expected failures are listed in KNOWN_FAILURES.md.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

anonymizer = pytest.importorskip(
    "anonymizer.anonymizer", reason="source repository or runtime dependencies unavailable"
)


class _SpyAnonymizer:
    """Test double: record module-level paths at construction and leave text unchanged."""

    seen: dict = {}

    def __init__(self, *args, **kwargs):
        _SpyAnonymizer.seen = {
            "LIST_PATH": str(anonymizer.LIST_PATH),
            "WHITELIST_PATH": str(anonymizer.WHITELIST_PATH),
        }

    def __call__(self, text):
        return text

    anonymize = __call__


@pytest.fixture
def patched_main(monkeypatch, tmp_path):
    inp = tmp_path / "in.xlsx"
    outp = tmp_path / "out.xlsx"
    custom_bl = tmp_path / "Mijn Blacklist.xlsx"
    custom_wl = tmp_path / "Mijn Whitelist.xlsx"
    for p in (custom_bl, custom_wl):
        p.write_bytes(b"")  # Exists but is not read by the test double.

    # Replace CombinedAnonymizer with the spy to avoid loading a model.
    monkeypatch.setattr(anonymizer, "CombinedAnonymizer", _SpyAnonymizer)
    # Return a small input frame from main()'s read_excel call.
    # This patches pandas globally during the test; read the output with openpyxl.
    monkeypatch.setattr(
        anonymizer.pd, "read_excel",
        lambda *a, **k: pd.DataFrame({"Toelichting_unmasked": ["hallo wereld"]}),
    )
    monkeypatch.setattr(
        sys, "argv",
        [
            "anonymizer", str(inp), str(outp),
            "--text-column", "Toelichting_unmasked",
            "--masked-text-column", "mijn_masker_kolom",
            "--blacklist-path", str(custom_bl),
            "--whitelist-path", str(custom_wl),
        ],
    )
    anonymizer.main()
    return {"out": outp, "custom_bl": str(custom_bl), "custom_wl": str(custom_wl)}


@pytest.mark.b4_redgreen  # V02 landed (2026-06-25): regression check.
def test_cli_blacklist_whitelist_reach_globals(patched_main):
    assert _SpyAnonymizer.seen["LIST_PATH"] == patched_main["custom_bl"]
    assert _SpyAnonymizer.seen["WHITELIST_PATH"] == patched_main["custom_wl"]


@pytest.mark.b4_redgreen  # V02 landed (2026-06-25): regression check.
def test_cli_masked_column_name_is_honored(patched_main):
    # Read the header with openpyxl because pandas.read_excel is patched.
    wb = load_workbook(patched_main["out"])
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "mijn_masker_kolom" in header
